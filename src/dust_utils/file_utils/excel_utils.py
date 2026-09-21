from pathlib import Path
from loguru import logger


class ExcelUtils:

    class WriteMode:
        # 覆盖写入
        OVERWRITE = "overwrite"
        # 追加下如
        APPEND = "append"

    @staticmethod
    def read(xlsx_path, sheet_name="Sheet1", cell_range=None):
        """读取 Excel，返回二维列表。

        Args:
            xlsx_path: Excel 文件路径。
            sheet_name: 工作表名称。
            cell_range: 读取范围，例如 "A1:F5"。
                        不指定时读取整个工作表。
        """
        from openpyxl import load_workbook
        from loguru import logger

        xlsx_path = Path(xlsx_path)

        if not xlsx_path.exists():
            logger.warning(f"文件不存在：{xlsx_path}")
            return None

        workbook = load_workbook(xlsx_path, data_only=False)

        try:
            if sheet_name not in workbook.sheetnames:
                logger.warning(
                    f"Excel 工作表不存在：{sheet_name}，"
                    f"当前工作表：{workbook.sheetnames}"
                )
                return None

            sheet = workbook[sheet_name]

            if cell_range:
                rows = sheet[cell_range]
            else:
                rows = sheet.iter_rows()

            result = [[cell.value for cell in row] for row in rows]

            logger.info(f"读取{len(result)}行数据：{xlsx_path}")

            return result

        finally:
            workbook.close()

    @staticmethod
    def write(
        xlsx_path,
        data,
        sheet_name="Sheet1",
        start_cell=None,
        write_mode=WriteMode.OVERWRITE,
    ):
        """写入 Excel。保留模板原有样式。"""
        from copy import copy
        from openpyxl import load_workbook, Workbook
        from openpyxl.utils.cell import coordinate_to_tuple

        xlsx_path = Path(xlsx_path)
        xlsx_path.parent.mkdir(parents=True, exist_ok=True)

        if write_mode not in ("overwrite", "append"):
            raise ValueError(f"不支持的写入模式：{write_mode}")

        # 打开或创建工作簿
        if xlsx_path.exists():
            workbook = load_workbook(xlsx_path)
        else:
            workbook = Workbook()

        # 获取工作表
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.create_sheet(sheet_name)

        # 确定起始位置
        if start_cell:
            row_index, col_index = coordinate_to_tuple(start_cell)
        elif write_mode == "append":
            row_index = sheet.max_row + 1
            col_index = 1
        else:
            row_index = 1
            col_index = 1

        # 写入数据
        for row_offset, row in enumerate(data):
            for col_offset, value in enumerate(row):
                target_row = row_index + row_offset
                target_col = col_index + col_offset

                cell = sheet.cell(
                    row=target_row,
                    column=target_col,
                )

                # 单元格本身没有样式时，继承列样式
                if cell.style_id == 0:
                    for dimension in sheet.column_dimensions.values():
                        if dimension.min <= target_col <= dimension.max:
                            if dimension.style:
                                cell._style = copy(
                                    workbook._cell_styles[dimension.style]
                                )
                            break

                cell.value = value

        logger.info(f"写入{len(data)}行数据：{xlsx_path}")

        workbook.save(xlsx_path)
        workbook.close()
